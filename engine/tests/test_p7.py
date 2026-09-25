"""0.10 기능: 점 집계, 비율·EB 보정, 시공간 분석."""

from __future__ import annotations

import shutil
from pathlib import Path

import geopandas as gpd
import libpysal
import numpy as np
import pytest
from conftest import make_grid
from shapely.geometry import Point


def _open(client, auth, path: Path) -> dict:
    r = client.post("/datasets/open", json={"path": str(path)}, headers=auth)
    assert r.status_code == 200, r.text
    return r.json()


def _weights(client, auth, ds_id: str, kind: str = "queen") -> str:
    r = client.post(f"/datasets/{ds_id}/weights", json={"type": kind}, headers=auth).json()
    return r["id"]


@pytest.fixture
def points_gpkg(tmp_path: Path) -> Path:
    """격자(EPSG:5186, 1km 칸) 위의 점. 경위도(4326)로 저장해 좌표계 변환도 확인함."""
    x0, y0 = 195_000.0, 545_000.0
    pts = [
        (x0 + 500, y0 + 500, 1.0),  # 칸 0
        (x0 + 600, y0 + 400, 3.0),  # 칸 0
        (x0 + 1500, y0 + 500, 10.0),  # 칸 1
        (x0 + 2000, y0 + 2500, 7.0),  # 칸 11·12 경계 → 한 칸에만 셈
        (x0 - 5000, y0, 99.0),  # 격자 밖
    ]
    gdf = gpd.GeoDataFrame(
        {"무게": [p[2] for p in pts]}, geometry=[Point(p[0], p[1]) for p in pts], crs=5186
    ).to_crs(4326)
    path = tmp_path / "points.gpkg"
    gdf.to_file(path, layer="points", engine="pyogrio")
    return path


def test_aggregate_points(client, auth, grid_gpkg, points_gpkg) -> None:
    grid = _open(client, auth, grid_gpkg)
    pts = _open(client, auth, points_gpkg)
    body = {
        "source_id": pts["id"],
        "count": True,
        "density": True,
        "stats": [{"column": "무게", "stat": "sum"}, {"column": "무게", "stat": "mean"}],
        "prefix": "PT",
    }
    r = client.post(f"/datasets/{grid['id']}/aggregate", json=body, headers=auth)
    assert r.status_code == 200, r.text
    res = r.json()
    outs = res["analysis"]["outputs"]
    assert outs == ["PT_CNT", "PT_DENS", "PT_SUM_무게", "PT_MEAN_무게"]
    assert any("밖에 있는" in n for n in res["notes"])
    assert any("좌표계" in n for n in res["notes"])

    rows = client.post(f"/datasets/{grid['id']}/rows", json={"limit": 50}, headers=auth).json()
    cols = rows["columns"]
    cnt = np.array([row[cols.index("PT_CNT")] for row in rows["rows"]])
    total = np.array([row[cols.index("PT_SUM_무게")] for row in rows["rows"]])
    mean = [row[cols.index("PT_MEAN_무게")] for row in rows["rows"]]
    dens = np.array([row[cols.index("PT_DENS")] for row in rows["rows"]])
    assert cnt.sum() == 4  # 격자 밖 1개 제외, 경계 위 점은 한 번만 셈
    assert cnt[0] == 2 and cnt[1] == 1
    assert total[0] == pytest.approx(4.0) and total[1] == pytest.approx(10.0)
    assert mean[0] == pytest.approx(2.0) and mean[5] is None  # 점 없는 칸의 평균은 값 없음
    assert total[5] == 0.0  # 점 없는 칸의 합계는 0
    assert dens[0] == pytest.approx(2.0, rel=1e-6)  # 1㎢ 칸에 2개

    # 같은 접두어로 다시 실행하면 대체함
    body["stats"] = []
    r = client.post(f"/datasets/{grid['id']}/aggregate", json=body, headers=auth).json()
    names = [c["name"] for c in r["info"]["columns"]]
    assert "PT_SUM_무게" not in names and "PT_CNT" in names

    bad = client.post(
        f"/datasets/{pts['id']}/aggregate", json={"source_id": grid["id"]}, headers=auth
    ).json()
    assert bad["error"]["code"] == "not_polygon"


def _sids(tmp_path: Path) -> tuple[Path, Path]:
    src = Path(libpysal.examples.get_path("sids2.shp")).parent
    dst = tmp_path / "sids2"
    shutil.copytree(src, dst)
    # 원본 GAL은 FIPSNO를 ID로 씀 → 행 순서(1..n) ID로 바꿔 저장함 (앱은 행 순서 ID만 읽음)
    w = libpysal.io.open(str(dst / "sids2.gal")).read()
    fips = gpd.read_file(dst / "sids2.shp")["FIPSNO"].astype(int).tolist()
    row = {f: i + 1 for i, f in enumerate(fips)}
    lines = [str(len(fips))]
    for f in fips:
        nb = (
            [row[int(k)] for k in w.neighbors[str(f)]]
            if str(f) in w.neighbors
            else [row[int(k)] for k in w.neighbors[f]]
        )
        lines += [f"{row[f]} {len(nb)}", " ".join(map(str, nb))]
    gal = dst / "sids2_rows.gal"
    gal.write_text("\n".join(lines) + "\n")
    return dst / "sids2.shp", gal


def test_rates_and_eb_moran(client, auth, tmp_path) -> None:
    from esda import smoothing

    shp, gal = _sids(tmp_path)
    ds = _open(client, auth, shp)
    wid = client.post(
        f"/datasets/{ds['id']}/weights/load", json={"path": str(gal)}, headers=auth
    ).json()["id"]
    g = gpd.read_file(shp)
    e, b = g["SID79"].to_numpy(float), g["BIR79"].to_numpy(float)

    url = f"/datasets/{ds['id']}/rates"
    for method, expected in (
        ("raw", e / b * 1000),
        ("eb", smoothing.Empirical_Bayes(e, b).r.ravel() * 1000),
        ("excess_risk", smoothing.Excess_Risk(e, b).r.ravel()),
    ):
        body = {"method": method, "event": "SID79", "base": "BIR79", "multiplier": 1000}
        r = client.post(url, json=body, headers=auth)
        assert r.status_code == 200, r.text
        col = r.json()["analysis"]["outputs"][0]
        vals = client.post(
            f"/datasets/{ds['id']}/columns", json={"names": [col]}, headers=auth
        ).content
        import pyarrow as pa

        got = pa.ipc.open_stream(vals).read_all().column(col).to_numpy()
        assert np.allclose(got, expected)

    need_w = client.post(url, json={"method": "spatial_eb", "event": "SID79", "base": "BIR79"})
    assert need_w.status_code in (401, 403)  # 인증 없이는 거부
    need_w = client.post(
        url, json={"method": "spatial_eb", "event": "SID79", "base": "BIR79"}, headers=auth
    ).json()
    assert need_w["error"]["code"] == "weights_required"
    r = client.post(
        url,
        json={"method": "spatial_eb", "event": "SID79", "base": "BIR79", "weights_id": wid},
        headers=auth,
    )
    assert r.status_code == 200, r.text

    # EB Moran's I: esda 문서·GeoDa 예제 값 (sids2, SID79/BIR79) = 0.1662
    m = client.post(
        f"/datasets/{ds['id']}/esda/moran",
        json={"column": "SID79", "rate_base": "BIR79", "weights_id": wid, "permutations": 99},
        headers=auth,
    ).json()
    assert m["I"] == pytest.approx(0.166223, abs=1e-5)
    assert "EB" in m["column"]

    loc = client.post(
        f"/datasets/{ds['id']}/esda/local",
        json={
            "method": "lisa_eb",
            "column": "SID79",
            "rate_base": "BIR79",
            "weights_id": wid,
            "permutations": 99,
        },
        headers=auth,
    ).json()
    assert loc["outputs"] == ["EBLISA_I", "EBLISA_CL", "EBLISA_P"]
    assert "EB 비율 LISA" in loc["description"]

    zero = client.post(url, json={"method": "raw", "event": "SID79", "base": "SID74"}, headers=auth)
    assert zero.json()["error"]["code"] == "nonpositive_base"


def _panel(tmp_path: Path) -> Path:
    """넓은 형태 시계열 격자: 값_2021~2023. 2023년에만 한쪽 모서리에 높은 값이 모임."""
    g = make_grid(6, 6)
    rng = np.random.default_rng(0)
    base = np.array([i % 6 + i // 6 for i in range(36)], dtype=float)
    g["값_2021"] = base + rng.normal(0, 0.1, 36)
    g["값_2022"] = base + rng.normal(0, 0.1, 36)
    g["값_2023"] = base[::-1] + rng.normal(0, 0.1, 36)
    path = tmp_path / "panel.gpkg"
    g.to_file(path, layer="panel", engine="pyogrio")
    return path


def test_time_groups_moran_lisa_and_project(client, auth, tmp_path) -> None:
    ds = _open(client, auth, _panel(tmp_path))
    did = ds["id"]
    guess = client.get(f"/datasets/{did}/time-groups/guess", headers=auth).json()
    assert guess == [
        {
            "name": "값",
            "columns": ["값_2021", "값_2022", "값_2023"],
            "labels": ["2021", "2022", "2023"],
        }
    ]
    info = client.put(f"/datasets/{did}/time-groups", json={"groups": guess}, headers=auth).json()
    assert info["time_groups"][0]["labels"] == ["2021", "2022", "2023"]
    bad = client.put(
        f"/datasets/{did}/time-groups",
        json={"groups": [{"name": "x", "columns": ["값_2021"], "labels": ["a"]}]},
        headers=auth,
    ).json()
    assert bad["error"]["code"] == "invalid_group"

    wid = _weights(client, auth, did, "rook")
    series = client.post(
        f"/datasets/{did}/timeseries/moran",
        json={"group": "값", "weights_id": wid, "permutations": 99},
        headers=auth,
    ).json()
    assert [r["label"] for r in series["rows"]] == ["2021", "2022", "2023"]
    assert all(r["I"] > 0.5 for r in series["rows"])  # 매끄러운 기울기 → 강한 양의 자기상관

    # 차분 Moran: 두 열의 차로 계산한 값 = 계산 필드로 만든 차의 Moran
    diff = client.post(
        f"/datasets/{did}/esda/moran",
        json={"column": "값_2023", "column_base": "값_2021", "weights_id": wid, "permutations": 0},
        headers=auth,
    ).json()
    client.post(
        f"/datasets/{did}/fields",
        json={"name": "차", "expression": "`값_2023` - `값_2021`"},
        headers=auth,
    )
    ref = client.post(
        f"/datasets/{did}/esda/moran",
        json={"column": "차", "weights_id": wid, "permutations": 0},
        headers=auth,
    ).json()
    assert diff["I"] == pytest.approx(ref["I"]) and diff["column"] == "값_2023−값_2021"

    r = client.post(
        f"/datasets/{did}/timeseries/lisa",
        json={"group": "값", "weights_id": wid, "permutations": 99, "prefix": "TL"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    a = r.json()["analysis"]
    assert a["outputs"] == ["TL_2021_CL", "TL_2022_CL", "TL_2023_CL", "TL_CHG"]
    rep = a["report"]
    assert rep["kind"] == "lisa_time" and len(rep["counts"]) == 3
    assert sum(map(sum, rep["matrix"])) == 36 * 2  # 인접 기간 2쌍 × 36개
    # 2022 → 2023에 값이 뒤집혀 High-High였던 곳이 Low-Low로 바뀐 곳이 있어야 함
    assert rep["matrix"][1][2] > 0 and rep["n_changed"] > 0

    # 프로젝트 저장·열기: 시간 변수 묶음과 기간별 LISA가 다시 만들어짐
    proj = tmp_path / "t.gstproj"
    client.post("/project/save", json={"path": str(proj), "datasets": [{"id": did}]}, headers=auth)
    opened = client.post("/project/open", json={"path": str(proj)}, headers=auth).json()
    info = opened["datasets"][0]["info"]
    assert info["time_groups"][0]["name"] == "값"
    names = [c["name"] for c in info["columns"]]
    assert "TL_CHG" in names and "차" in names
    assert opened["datasets"][0]["warnings"] == []


def test_pivot_long_and_merge_files(client, auth, tmp_path) -> None:
    # 긴 형태: 격자 20칸 × 3년 = 60행 (같은 지오메트리가 해마다 반복됨)
    g = make_grid()
    rows = []
    for year in (2022, 2023, 2024):
        part = g.copy()
        part["연도"] = year
        part["어획량"] = part["값"] * (year - 2021)
        rows.append(part)
    long = gpd.GeoDataFrame(np.concatenate([p.to_numpy() for p in rows]), columns=rows[0].columns)
    long = gpd.GeoDataFrame(long, geometry="geometry", crs=5186)
    long["어획량"] = long["어획량"].astype(float)
    long["연도"] = long["연도"].astype(int)
    long_path = tmp_path / "long.gpkg"
    long.to_file(long_path, layer="long", engine="pyogrio")
    ds = _open(client, auth, long_path)
    r = client.post(
        "/timeseries/pivot",
        json={
            "dataset_id": ds["id"],
            "id_column": "이름",
            "time_column": "연도",
            "value_columns": ["어획량"],
            "path": str(tmp_path / "wide"),
        },
        headers=auth,
    )
    assert r.status_code == 200, r.text
    info = r.json()["info"]
    assert info["n_rows"] == 20 and info["path"].endswith("wide.gpkg")
    assert info["time_groups"] == [
        {
            "name": "어획량",
            "columns": ["어획량_2022", "어획량_2023", "어획량_2024"],
            "labels": ["2022", "2023", "2024"],
        }
    ]

    # 기간별 파일: 2개 파일을 이름(ID)으로 이어 붙임. 두 번째 파일은 행 순서가 다르고 한 칸이 빠짐
    a = make_grid()
    b = make_grid().iloc[::-1].iloc[1:].copy()
    b["값"] = b["값"] * 10
    pa_, pb = tmp_path / "y2023.gpkg", tmp_path / "y2024.gpkg"
    a.to_file(pa_, layer="a", engine="pyogrio")
    b.to_file(pb, layer="b", engine="pyogrio")
    da, db = _open(client, auth, pa_), _open(client, auth, pb)
    r = client.post(
        "/timeseries/merge",
        json={
            "parts": [
                {"dataset_id": da["id"], "id_column": "이름", "label": "2023"},
                {"dataset_id": db["id"], "id_column": "이름", "label": "2024"},
            ],
            "value_columns": ["값"],
            "path": str(tmp_path / "merged.gpkg"),
        },
        headers=auth,
    ).json()
    assert r["info"]["time_groups"][0]["columns"] == ["값_2023", "값_2024"]
    assert any("값 없음" in n for n in r["notes"])
    m = gpd.read_file(tmp_path / "merged.gpkg")
    first = m.set_index("ID").loc["격자_0101"]
    assert first["값_2024"] == pytest.approx(first["값_2023"] * 10)
    assert m["값_2024"].isna().sum() == 1


def test_aggregate_project_cache(client, auth, grid_gpkg, points_gpkg, tmp_path) -> None:
    grid = _open(client, auth, grid_gpkg)
    pts = _open(client, auth, points_gpkg)
    client.post(
        f"/datasets/{grid['id']}/aggregate",
        json={"source_id": pts["id"], "stats": [{"column": "무게", "stat": "max"}]},
        headers=auth,
    )
    client.post(
        f"/datasets/{grid['id']}/fields",
        json={"name": "분모", "expression": "`정수` + 1"},
        headers=auth,
    )
    r = client.post(
        f"/datasets/{grid['id']}/rates",
        json={"method": "raw", "event": "PT_CNT", "base": "분모", "name": "비율"},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    proj = tmp_path / "agg.gstproj"
    client.post(
        "/project/save",
        json={"path": str(proj), "datasets": [{"id": grid["id"]}, {"id": pts["id"]}]},
        headers=auth,
    )
    assert (tmp_path / "agg.gstcache" / "dataset0.parquet").exists()
    opened = client.post("/project/open", json={"path": str(proj)}, headers=auth).json()
    names = [c["name"] for c in opened["datasets"][0]["info"]["columns"]]
    assert "PT_MAX_무게" in names and "비율" in names
    assert opened["datasets"][0]["warnings"] == []
    # 앱에서 만든 가중치로 공간 비율·공간 EB (esda가 요구하는 id_order가 없는 W)
    g2 = opened["datasets"][0]["info"]["id"]
    wid = _weights(client, auth, g2, "rook")
    for method in ("spatial_rate", "spatial_eb"):
        r = client.post(
            f"/datasets/{g2}/rates",
            json={"method": method, "event": "PT_CNT", "base": "분모", "weights_id": wid},
            headers=auth,
        )
        assert r.status_code == 200, r.text
    # 캐시가 없어도 원본 점 파일을 다시 읽어 계산함
    shutil.rmtree(tmp_path / "agg.gstcache")
    opened = client.post("/project/open", json={"path": str(proj)}, headers=auth).json()
    first = opened["datasets"][0]
    assert "PT_CNT" in [c["name"] for c in first["info"]["columns"]]
    assert any("다시 읽어" in w for w in first["warnings"])


def test_pooled_classification(client, auth, tmp_path) -> None:
    ds = _open(client, auth, _panel(tmp_path))
    url = f"/datasets/{ds['id']}/classify"
    cols = ["값_2021", "값_2022", "값_2023"]
    own = client.post(
        url, json={"column": "값_2021", "method": "equal_interval", "k": 4}, headers=auth
    )
    pooled = client.post(
        url,
        json={"column": "값_2021", "method": "equal_interval", "k": 4, "pool": cols},
        headers=auth,
    ).json()
    other = client.post(
        url,
        json={"column": "값_2023", "method": "equal_interval", "k": 4, "pool": cols},
        headers=auth,
    ).json()
    assert pooled["breaks"] == other["breaks"]  # 모든 기간이 같은 경계를 씀
    assert own.status_code == 200 and sum(pooled["counts"]) == 36
    box = client.post(
        url, json={"column": "값_2022", "method": "box_plot", "pool": cols}, headers=auth
    ).json()
    assert len(box["labels"]) == 6


def test_report_export(client, auth, tmp_path, grid_gpkg, points_gpkg) -> None:
    from docx import Document

    ds = _open(client, auth, _panel(tmp_path))
    did = ds["id"]
    guess = client.get(f"/datasets/{did}/time-groups/guess", headers=auth).json()
    client.put(f"/datasets/{did}/time-groups", json={"groups": guess}, headers=auth)
    wid = _weights(client, auth, did)
    client.post(
        f"/datasets/{did}/esda/local",
        json={"method": "lisa", "column": "값_2021", "weights_id": wid, "permutations": 99},
        headers=auth,
    )
    client.post(
        f"/datasets/{did}/timeseries/lisa",
        json={"group": "값", "weights_id": wid, "permutations": 99},
        headers=auth,
    )
    grid = _open(client, auth, grid_gpkg)
    pts = _open(client, auth, points_gpkg)
    client.post(f"/datasets/{grid['id']}/aggregate", json={"source_id": pts["id"]}, headers=auth)

    from geostat_engine.io.raster import encode_png

    png = encode_png(np.full((8, 8, 4), 200, dtype=np.uint8))
    import base64

    body = {
        "path": str(tmp_path / "보고서"),
        "format": "html",
        "title": "시험 보고서",
        "datasets": [{"dataset_id": did}, {"dataset_id": grid["id"]}],
        "images": [{"data": base64.b64encode(png).decode(), "caption": "지도"}],
        "morans": [{"dataset": "panel", "column": "값_2021", "weights": "queen1", "I": 0.61}],
    }
    r = client.post("/report/export", json=body, headers=auth)
    assert r.status_code == 200, r.text
    text = Path(r.json()["path"]).read_text(encoding="utf-8")
    assert r.json()["path"].endswith("보고서.html")
    for needle in (
        "시험 보고서",
        "시간 변수 묶음",
        "군집 전이표",
        "High-High",
        "점 집계",
        "data:image/png",
    ):
        assert needle in text, needle

    body["format"] = "docx"
    r = client.post("/report/export", json=body, headers=auth)
    assert r.status_code == 200, r.text
    doc = Document(r.json()["path"])
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert "시험 보고서" in heads and "panel" in heads
    assert len(doc.tables) >= 6 and len(doc.inline_shapes) == 1

    body["images"] = [{"data": base64.b64encode(b"nope").decode()}]
    bad = client.post("/report/export", json=body, headers=auth).json()
    assert bad["error"]["code"] == "invalid_image"


def test_report_blocks_regression_cluster(client, auth, tmp_path) -> None:
    from geostat_engine import report, service

    ds_info = _open(client, auth, _panel(tmp_path))
    wid = _weights(client, auth, ds_info["id"])
    ds = client.app.state.geostat.get(ds_info["id"])
    service.run_regression_sync(
        ds, {"model": "ols", "y": "값_2022", "x": ["값_2021"], "weights_id": wid}
    )
    service.run_regression_sync(
        ds, {"model": "gwr", "y": "값_2022", "x": ["값_2021"], "weights_id": wid}
    )
    service.run_cluster_sync(
        ds,
        {
            "method": "skater",
            "variables": ["값_2021", "값_2022"],
            "n_clusters": 3,
            "weights_id": wid,
        },
    )
    blocks = report.build("t", [(ds, None)], [])
    html_text = report.render_html(blocks)
    assert "계수" in html_text and "지역 계수 요약" in html_text and "군집 내 SS" in html_text
    assert report.render_docx(blocks)[:2] == b"PK"
