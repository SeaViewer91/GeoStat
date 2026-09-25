"""분석 보고서 내보내기 (HTML·Word).

데이터셋마다 자료 정보, 공간가중치, 시간 변수 묶음, 분석 기록(회귀·군집·국지 통계·점 집계 등)을 모아
블록 목록(제목·문단·표·그림)으로 만든 뒤 HTML이나 docx로 그림. 지도 그림은 앱이 PNG로 보내 줌.
용역 보고서 부록처럼 그대로 붙이거나 Word에서 고쳐 쓰는 용도임.
"""

from __future__ import annotations

import base64
import html
import io
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from geostat_engine import __version__
from geostat_engine.analysis.regression import _fmt
from geostat_engine.errors import EngineError
from geostat_engine.state import AnalysisRecord, Dataset

# 국지 통계 군집 코드 이름 (esda_ops의 코드와 같음)
_CLUSTER_NAMES = {
    "lisa": ["유의하지 않음", "High-High", "Low-Low", "Low-High", "High-Low", "이웃 없음"],
    "gi_star": ["유의하지 않음", "핫스팟", "콜드스팟", "", "", "이웃 없음"],
    "local_geary": [
        "유의하지 않음",
        "High-High",
        "Low-Low",
        "기타 양의 연관",
        "음의 연관",
        "이웃 없음",
    ],
}
_CLUSTER_NAMES["lisa_bv"] = _CLUSTER_NAMES["lisa_eb"] = _CLUSTER_NAMES["lisa"]


@dataclass
class Block:
    kind: str  # h1 h2 h3 p note table image
    text: str = ""
    header: list[str] = field(default_factory=list)
    rows: list[list[Any]] = field(default_factory=list)
    image: bytes | None = None


def _h(level: int, text: str) -> Block:
    return Block(kind=f"h{level}", text=text)


def _p(text: str) -> Block:
    return Block(kind="p", text=text)


def _table(header: list[str], rows: list[list[Any]]) -> Block:
    return Block(kind="table", header=header, rows=[[_fmt(v) for v in r] for r in rows])


def _kv(pairs: list[tuple[str, Any]]) -> Block:
    return _table(["항목", "값"], [[k, v] for k, v in pairs])


# ---- 내용 만들기 -------------------------------------------------------------------


def build(
    title: str,
    datasets: list[tuple[Dataset, list[str] | None]],
    images: list[tuple[bytes, str]],
    morans: list[dict[str, Any]] | None = None,
) -> list[Block]:
    blocks = [
        _h(1, title or "GeoStat 분석 보고서"),
        Block(
            kind="note",
            text=f"GeoStat {__version__} · {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M')}",
        ),
    ]
    for png, caption in images:
        blocks.append(Block(kind="image", image=png, text=caption))
    for ds, ids in datasets:
        blocks += _dataset_blocks(ds, ids)
    if morans:
        blocks.append(_h(2, "전역 Moran's I (차트 패널)"))
        blocks.append(
            _table(
                ["자료", "변수", "가중치", "I", "E[I]", "z (순열)", "유사 p값", "순열"],
                [
                    [
                        m.get("dataset", ""),
                        m.get("column", ""),
                        m.get("weights", ""),
                        m.get("I"),
                        m.get("expected"),
                        m.get("z_sim"),
                        m.get("p_sim"),
                        m.get("permutations"),
                    ]
                    for m in morans
                ],
            )
        )
    return blocks


def _dataset_blocks(ds: Dataset, ids: list[str] | None) -> list[Block]:
    gdf = ds.gdf
    crs = gdf.crs
    out = [_h(2, ds.name)]
    geom_kinds = ", ".join(sorted(set(gdf.geom_type.dropna()))) or "-"
    out.append(
        _kv(
            [
                ("원본", str(ds.path)),
                ("피처 수", len(gdf)),
                ("지오메트리", geom_kinds),
                ("좌표계", f"EPSG:{crs.to_epsg()} {crs.name}" if crs else "없음"),
                ("속성 인코딩", ds.encoding or "-"),
            ]
        )
    )
    if ds.fields:
        out.append(_h(3, "계산 필드"))
        out.append(_table(["열", "식"], [[f.name, f.expression] for f in ds.fields]))
    if ds.time_groups:
        out.append(_h(3, "시간 변수 묶음"))
        out.append(
            _table(
                ["묶음", "기간", "열"],
                [[g.name, ", ".join(g.labels), ", ".join(g.columns)] for g in ds.time_groups],
            )
        )
    if ds.weights:
        out.append(_h(3, "공간가중치"))
        rows = []
        for w in ds.weights.values():
            s = w.summary
            rows.append(
                [
                    w.name,
                    w.description,
                    s.get("mean_neighbors"),
                    s.get("min_neighbors"),
                    s.get("max_neighbors"),
                    s.get("n_islands"),
                ]
            )
        out.append(_table(["이름", "방법", "평균 이웃", "최소", "최대", "이웃 없음"], rows))
    analyses = [a for a in ds.analyses if ids is None or a.id in ids]
    if analyses:
        out.append(_h(3, "분석"))
    for a in analyses:
        out += _analysis_blocks(a)
    return out


def _analysis_blocks(a: AnalysisRecord) -> list[Block]:
    out = [Block(kind="h4", text=a.description)]
    report = a.summary.get("report")
    if a.method in _CLUSTER_NAMES:
        names = _CLUSTER_NAMES[a.method]
        counts = a.summary.get("counts", {})
        rows = [
            [names[int(c)] or str(c), n] for c, n in sorted(counts.items(), key=lambda x: int(x[0]))
        ]
        out.append(_table(["구분", "피처 수"], rows))
        if a.summary.get("threshold") is not None:
            out.append(_p(f"유의 판정 기준 p ≤ {_fmt(a.summary['threshold'], 4)}"))
    elif a.method == "regression" and report:
        out += _regression_blocks(report)
    elif a.method == "cluster" and report:
        out += _cluster_blocks(report)
    elif a.method == "lisa_time" and report:
        out += _lisa_time_blocks(report)
    elif a.method == "aggregate" and a.summary:
        s = a.summary
        if "n_source" in s:
            out.append(
                _kv(
                    [
                        ("원본 피처", s["n_source"]),
                        ("집계된 피처", s["n_matched"]),
                        ("점이 없는 폴리곤", s["n_empty"]),
                    ]
                )
            )
        out += [Block(kind="note", text=n) for n in s.get("notes", [])]
    elif a.method == "rate" and a.summary:
        s = a.summary
        out.append(
            _kv([("최솟값", s.get("min")), ("평균", s.get("mean")), ("최댓값", s.get("max"))])
        )
    out.append(Block(kind="note", text="결과 열: " + ", ".join(a.outputs)))
    return out


def _regression_blocks(r: dict[str, Any]) -> list[Block]:
    out = [
        _p(f"종속변수: {r['y']} · 독립변수: {', '.join(r['x'])} · 관측치 {r['n']:,}"),
        _kv([(k, v) for k, v in r["summary"]]),
    ]
    if r.get("coefficients"):
        stat = r.get("stat_label", "t")
        out.append(
            _table(
                ["변수", "계수", "표준오차", f"{stat}값", "유의확률"],
                [[c["name"], c["coef"], c["se"], c["stat"], c["p"]] for c in r["coefficients"]],
            )
        )
    if r.get("impacts"):
        imp = r["impacts"]
        out.append(_p(f"직접·간접·총 효과 ({imp['method']})"))
        out.append(
            _table(
                ["변수", "직접", "간접", "총"],
                [[x["name"], x["direct"], x["indirect"], x["total"]] for x in imp["rows"]],
            )
        )
    if r.get("local"):
        out.append(_p("지역 계수 요약"))
        out.append(
            _table(
                ["변수", "대역폭", "평균", "최소", "중앙값", "최대", "유의 %"],
                [
                    [
                        x["name"],
                        x["bandwidth"],
                        x["mean"],
                        x["min"],
                        x["median"],
                        x["max"],
                        x.get("pct_significant"),
                    ]
                    for x in r["local"]
                ],
            )
        )
    if r.get("diagnostics"):
        out.append(_p("진단"))
        out.append(
            _table(
                ["구분", "검정", "값", "자유도", "유의확률"],
                [[d["group"], d["name"], d["value"], d["df"], d["p"]] for d in r["diagnostics"]],
            )
        )
    out += [Block(kind="note", text=n) for n in r.get("notes", []) if n]
    return out


def _cluster_blocks(r: dict[str, Any]) -> list[Block]:
    out = [_kv([(k, v) for k, v in r["summary"]])]
    header = ["군집", "크기", "군집 내 SS", "공간 조각"] + [f"{v} 평균" for v in r["variables"]]
    rows = [
        [c["id"], c["size"], c.get("within_ss"), c.get("fragments"), *c.get("means", [])]
        for c in r["clusters"]
    ]
    out.append(_table(header, rows))
    out += [Block(kind="note", text=n) for n in r.get("notes", []) if n]
    return out


def _lisa_time_blocks(r: dict[str, Any]) -> list[Block]:
    cats = r["categories"]
    shown = [i for i in range(len(cats)) if any(c[i] for c in r["counts"])]
    out = [
        _p("기간별 군집 수"),
        _table(
            ["기간"] + [cats[i] for i in shown],
            [[lb] + [c[i] for i in shown] for lb, c in zip(r["labels"], r["counts"], strict=True)],
        ),
        _p("군집 전이표 (행: 앞 기간 → 열: 다음 기간, 인접 기간 합계)"),
        _table(
            ["앞 \\ 뒤"] + [cats[i] for i in shown],
            [[cats[i]] + [r["matrix"][i][j] for j in shown] for i in shown],
        ),
        _kv(
            [
                ("군집이 한 번 이상 바뀐 피처", r["n_changed"]),
                ("모든 기간 같은 유의 군집인 피처", r["n_stable_cluster"]),
            ]
        ),
    ]
    return out


def lisa_time_text(r: dict[str, Any], description: str = "") -> str:
    """기간별 LISA 텍스트 보고서 (결과 패널의 '보고서 저장')."""
    cats = r["categories"]
    lines = ["=" * 72, f"GeoStat 시공간 분석 보고서 — {r['title']}", "=" * 72]
    if description:
        lines.append(f"설정: {description}")
    lines += ["", "기간별 군집 수", "-" * 72]
    lines.append(f"{'기간':<12}" + "".join(f"{c:>12}" for c in cats))
    for lb, row in zip(r["labels"], r["counts"], strict=True):
        lines.append(f"{lb:<12}" + "".join(f"{v:>12,}" for v in row))
    lines += ["", "군집 전이표 (행: 앞 기간 → 열: 다음 기간)", "-" * 72]
    lines.append(f"{'':<12}" + "".join(f"{c:>12}" for c in cats))
    for c, row in zip(cats, r["matrix"], strict=True):
        lines.append(f"{c:<12}" + "".join(f"{v:>12,}" for v in row))
    lines += [
        "",
        f"군집이 한 번 이상 바뀐 피처: {r['n_changed']:,}",
        f"모든 기간 같은 유의 군집인 피처: {r['n_stable_cluster']:,}",
        "",
    ]
    return "\n".join(lines)


# ---- 그리기 --------------------------------------------------------------------

_CSS = """
body{font-family:-apple-system,"Apple SD Gothic Neo","Malgun Gothic","Noto Sans KR",sans-serif;
max-width:960px;margin:32px auto;padding:0 20px;color:#1d2433;line-height:1.55}
h1{font-size:26px;border-bottom:2px solid #2b6cb0;padding-bottom:6px}
h2{font-size:20px;margin-top:36px;border-bottom:1px solid #d5dbe5;padding-bottom:4px}
h3{font-size:16px;margin-top:24px;color:#2b4a7a}
h4{font-size:14px;margin:18px 0 6px}
table{border-collapse:collapse;margin:8px 0 12px;font-size:13px}
th,td{border:1px solid #cfd6e2;padding:4px 8px;text-align:right;vertical-align:top}
th{background:#eef2f8;text-align:center}
td:first-child{text-align:left}
.note{color:#5b6475;font-size:12px;margin:4px 0}
figure{margin:12px 0}figure img{max-width:100%;border:1px solid #d5dbe5}
figcaption{color:#5b6475;font-size:12px}
@media print{body{margin:0;max-width:none}}
"""


def render_html(blocks: list[Block]) -> str:
    esc = html.escape
    title = next((b.text for b in blocks if b.kind == "h1"), "GeoStat")
    parts = [
        "<!doctype html><html lang='ko'><head><meta charset='utf-8'>",
        f"<title>{esc(title)}</title><style>{_CSS}</style></head><body>",
    ]
    for b in blocks:
        if b.kind in ("h1", "h2", "h3", "h4"):
            parts.append(f"<{b.kind}>{esc(b.text)}</{b.kind}>")
        elif b.kind == "p":
            parts.append(f"<p>{esc(b.text)}</p>")
        elif b.kind == "note":
            parts.append(f"<p class='note'>{esc(b.text)}</p>")
        elif b.kind == "table":
            head = "".join(f"<th>{esc(str(h))}</th>" for h in b.header)
            body = "".join(
                "<tr>" + "".join(f"<td>{esc(str(v))}</td>" for v in row) + "</tr>" for row in b.rows
            )
            parts.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        elif b.kind == "image" and b.image:
            data = base64.b64encode(b.image).decode("ascii")
            cap = f"<figcaption>{esc(b.text)}</figcaption>" if b.text else ""
            parts.append(f"<figure><img src='data:image/png;base64,{data}' alt=''>{cap}</figure>")
    parts.append("</body></html>")
    return "".join(parts)


def render_docx(blocks: list[Block]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    # 한글 글꼴: 동아시아 글꼴을 따로 지정해야 Word에서 한글이 제대로 나옴
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[style_name]
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        fonts.set(qn("w:eastAsia"), "맑은 고딕")
    doc.styles["Normal"].font.size = Pt(10)

    for b in blocks:
        if b.kind == "h1":
            doc.add_heading(b.text, level=0)
        elif b.kind in ("h2", "h3", "h4"):
            doc.add_heading(b.text, level=int(b.kind[1]) - 1)
        elif b.kind == "p":
            doc.add_paragraph(b.text)
        elif b.kind == "note":
            run = doc.add_paragraph().add_run(b.text)
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x5B, 0x64, 0x75)
        elif b.kind == "table":
            table = doc.add_table(rows=1, cols=len(b.header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, text in zip(table.rows[0].cells, b.header, strict=True):
                cell.text = str(text)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
            for row in b.rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=True):
                    cell.text = str(text)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
            doc.add_paragraph()
        elif b.kind == "image" and b.image:
            doc.add_picture(io.BytesIO(b.image), width=Cm(16))
            if b.text:
                doc.add_paragraph(b.text).runs[0].font.size = Pt(8.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def write(path: Path, fmt: str, blocks: list[Block]) -> Path:
    suffix = ".docx" if fmt == "docx" else ".html"
    if path.suffix.lower() != suffix:
        path = path.with_suffix(suffix)
    if not path.parent.exists():
        raise EngineError("folder_not_found", f"폴더가 없음: {path.parent}")
    try:
        if fmt == "docx":
            path.write_bytes(render_docx(blocks))
        else:
            path.write_text(render_html(blocks), encoding="utf-8")
    except OSError as exc:
        raise EngineError("write_failed", f"보고서를 저장하지 못함: {exc}") from exc
    return path
